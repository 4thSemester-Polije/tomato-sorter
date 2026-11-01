"""Train and evaluate separate ResNeXt-101 ripeness and defect classifiers."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

import numpy as np
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, f1_score
from sklearn.model_selection import train_test_split

from ml.resnext import (
    DEFECT_CLASSES,
    DEFECT_LEVELS,
    INPUT_SIZE,
    MODEL_NAME,
    RIPENESS_CLASSES,
    RIPENESS_LEVELS,
    create_model,
    get_preprocess_input,
    load_trusted_model,
)
from ml.rubric import grade_from_levels

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class Sample:
    path: Path
    label: str
    digest: str


def _collect_samples(class_directories: dict[str, Path]) -> list[Sample]:
    """Deduplicate identical files and exclude copies carrying conflicting labels."""
    by_digest: dict[str, list[Sample]] = {}
    for label, directory in class_directories.items():
        if not directory.is_dir():
            raise FileNotFoundError(f"Dataset class directory not found: {directory}")
        for path in sorted(directory.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in {
                ".jpg",
                ".jpeg",
                ".png",
            }:
                continue
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            by_digest.setdefault(digest, []).append(Sample(path, label, digest))

    if not by_digest:
        raise ValueError(f"No supported JPG/PNG images found in {class_directories}.")

    unique_samples: list[Sample] = []
    conflicting_groups: list[list[Sample]] = []
    for copies in by_digest.values():
        labels = {sample.label for sample in copies}
        if len(labels) > 1:
            conflicting_groups.append(copies)
        else:
            unique_samples.append(copies[0])

    if conflicting_groups:
        logger.warning(
            "Excluded %d identical-image group(s) with conflicting labels (%d files): %s",
            len(conflicting_groups),
            sum(len(group) for group in conflicting_groups),
            [
                [str(sample.path) for sample in group]
                for group in conflicting_groups
            ],
        )
    return sorted(unique_samples, key=lambda sample: str(sample.path))


def _split_samples(
    samples: Sequence[Sample],
    seed: int,
) -> tuple[list[Sample], list[Sample], list[Sample]]:
    """Stratify by class into 70/15/15 partitions after exact-image deduplication."""
    labels = [sample.label for sample in samples]
    counts = Counter(labels)
    if len(counts) < 2 or min(counts.values()) < 3:
        raise ValueError(
            "Each class needs at least three unique images for a stratified split; "
            f"found {dict(counts)}."
        )
    train_validation, test = train_test_split(
        list(samples),
        test_size=0.15,
        random_state=seed,
        stratify=labels,
    )
    train, validation = train_test_split(
        train_validation,
        test_size=0.15 / 0.85,
        random_state=seed,
        stratify=[sample.label for sample in train_validation],
    )
    return list(train), list(validation), list(test)


def _load_tensorflow_dataset(
    samples: Sequence[Sample],
    class_to_index: dict[str, int],
    preprocess_input: Any,
    batch_size: int,
    training: bool,
) -> Any:
    import tensorflow as tf

    dataset = tf.data.Dataset.from_tensor_slices(
        ([str(sample.path) for sample in samples],
         [class_to_index[sample.label] for sample in samples])
    )

    def load_image(path: Any, label: Any) -> tuple[Any, Any]:
        image = tf.io.decode_image(
            tf.io.read_file(path),
            channels=3,
            expand_animations=False,
        )
        image.set_shape((None, None, 3))
        image = tf.image.resize(image, INPUT_SIZE)
        return preprocess_input(image), label

    dataset = dataset.map(load_image, num_parallel_calls=tf.data.AUTOTUNE)
    if training:
        dataset = dataset.shuffle(
            len(samples),
            seed=42,
            reshuffle_each_iteration=True,
        )
    return dataset.batch(batch_size).prefetch(tf.data.AUTOTUNE)


def _classification_metrics(
    expected: Sequence[int],
    predicted: Sequence[int],
    class_names: Sequence[str],
) -> dict[str, Any]:
    from sklearn.metrics import precision_recall_fscore_support

    labels = list(range(len(class_names)))
    precision, recall, f1, _ = precision_recall_fscore_support(
        expected, predicted, labels=labels, zero_division=0
    )
    return {
        "accuracy": float(accuracy_score(expected, predicted)),
        "macro_f1": float(f1_score(expected, predicted, average="macro")),
        "classes": list(class_names),
        "precision_per_class": {
            name: float(precision[index]) for index, name in enumerate(class_names)
        },
        "recall_per_class": {
            name: float(recall[index]) for index, name in enumerate(class_names)
        },
        "f1_per_class": {
            name: float(f1[index]) for index, name in enumerate(class_names)
        },
        "confusion_matrix": confusion_matrix(
            expected, predicted, labels=labels
        ).tolist(),
        "classification_report": classification_report(
            expected,
            predicted,
            labels=labels,
            target_names=list(class_names),
            output_dict=True,
            zero_division=0,
        ),
    }


def _metric_report(
    model: Any,
    dataset: Any,
    class_names: Sequence[str],
) -> dict[str, Any]:
    expected = np.concatenate([labels.numpy() for _, labels in dataset])
    predicted = np.argmax(model.predict(dataset, verbose=0), axis=1)
    return _classification_metrics(expected, predicted, class_names)


def _train_task(
    task: str,
    samples: list[Sample],
    class_names: Sequence[str],
    class_levels: dict[str, int],
    output_path: Path,
    seed: int,
    batch_size: int,
    head_epochs: int,
    fine_tune_epochs: int,
    reuse_existing: bool = False,
) -> tuple[Any, dict[str, Any]]:
    import tensorflow as tf

    train_samples, validation_samples, test_samples = _split_samples(samples, seed)
    class_to_index = {name: index for index, name in enumerate(class_names)}
    output_path.parent.mkdir(parents=True, exist_ok=True)
    checkpoint_reused = reuse_existing and output_path.is_file()
    if checkpoint_reused:
        logger.info("Reusing existing model checkpoint: %s", output_path)
        model = load_trusted_model(output_path)
        backbone = None
        preprocess_input = get_preprocess_input()
        history: dict[str, list[float]] = {}
    else:
        model, backbone, preprocess_input = create_model(task, len(class_names))
        model.compile(
            optimizer=tf.keras.optimizers.Adam(learning_rate=1e-3),
            loss="sparse_categorical_crossentropy",
            metrics=["accuracy"],
        )
        history = {}
    train_data = _load_tensorflow_dataset(
        train_samples, class_to_index, preprocess_input, batch_size, True
    )
    validation_data = _load_tensorflow_dataset(
        validation_samples, class_to_index, preprocess_input, batch_size, False
    )
    test_data = _load_tensorflow_dataset(
        test_samples, class_to_index, preprocess_input, batch_size, False
    )

    if not checkpoint_reused:
        callbacks = [
            tf.keras.callbacks.EarlyStopping(
                monitor="val_loss", patience=3, restore_best_weights=True
            ),
            tf.keras.callbacks.ModelCheckpoint(
                filepath=str(output_path),
                monitor="val_loss",
                save_best_only=True,
            ),
        ]
        fit_history = model.fit(
            train_data,
            validation_data=validation_data,
            epochs=head_epochs,
            callbacks=callbacks,
            verbose=2,
        )
        history.update(
            {key: [float(value) for value in values] for key, values in fit_history.history.items()}
        )
        if fine_tune_epochs > 0:
            backbone.trainable = True
            for layer in backbone.layers[:-30]:
                layer.trainable = False
            for layer in backbone.layers:
                if isinstance(layer, tf.keras.layers.BatchNormalization):
                    layer.trainable = False
            model.compile(
                optimizer=tf.keras.optimizers.Adam(learning_rate=1e-5),
                loss="sparse_categorical_crossentropy",
                metrics=["accuracy"],
            )
            fine_tune_history = model.fit(
                train_data,
                validation_data=validation_data,
                epochs=fine_tune_epochs,
                callbacks=callbacks,
                verbose=2,
            )
            history.update(
                {
                    f"fine_tune_{key}": [float(value) for value in values]
                    for key, values in fine_tune_history.history.items()
                }
            )

    metadata = {
        "model": MODEL_NAME,
        "task": task,
        "input_size": list(INPUT_SIZE),
        "class_names": list(class_names),
        "class_levels": class_levels,
        "pretrained_weights": "ImageNet",
        "seed": seed,
        "split_percentages": {"train": 70, "validation": 15, "test": 15},
        "unique_images": {
            "total": len(samples),
            "train": len(train_samples),
            "validation": len(validation_samples),
            "test": len(test_samples),
        },
        "split_manifest": {
            "train": [
                {"sha256": sample.digest, "class": sample.label, "file": sample.path.name}
                for sample in train_samples
            ],
            "validation": [
                {
                    "sha256": sample.digest,
                    "class": sample.label,
                    "file": sample.path.name,
                }
                for sample in validation_samples
            ],
            "test": [
                {"sha256": sample.digest, "class": sample.label, "file": sample.path.name}
                for sample in test_samples
            ],
        },
        "test_metrics": _metric_report(model, test_data, class_names),
        "history": history,
        "checkpoint_reused": checkpoint_reused,
    }
    return model, metadata


def _primary_test_samples(data_root: Path) -> list[tuple[Path, dict[str, Any]]]:
    try:
        import pillow_heif
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Primary HEIC evaluation requires pillow-heif and openpyxl."
        ) from exc

    pillow_heif.register_heif_opener()
    primary_root = data_root / "Primer"
    photo_directory = primary_root / "Foto"
    labels_path = primary_root / "Label" / "Labelv.1.xlsx"
    mapping_path = primary_root / "Label" / "mapping_nama_file.csv"
    for required_path in (photo_directory, labels_path, mapping_path):
        if not required_path.exists():
            raise FileNotFoundError(f"Primary evaluation input not found: {required_path}")

    with mapping_path.open(newline="", encoding="utf-8-sig") as file:
        filename_by_id = {
            Path(row["nama_file_baru"]).stem: row["nama_file_baru"]
            for row in csv.DictReader(file)
        }

    worksheet = load_workbook(labels_path, read_only=True, data_only=True)["Pelabelan"]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    samples: list[tuple[Path, dict[str, Any]]] = []
    for values in rows:
        row = dict(zip(headers, values))
        tomato_id = row.get("id_tomat")
        if not tomato_id:
            continue
        filename = filename_by_id.get(str(tomato_id))
        if filename is None:
            raise ValueError(f"No mapped photo filename for tomato {tomato_id}.")
        path = photo_directory / filename
        if not path.is_file():
            raise FileNotFoundError(f"Primary photo not found: {path}")
        for column in ("tingkat_kematangan", "tingkat_cacat", "grade"):
            if row.get(column) is None:
                raise ValueError(f"Missing {column} label for tomato {tomato_id}.")
        ripeness_level = int(row["tingkat_kematangan"])
        defect_level = int(row["tingkat_cacat"])
        if ripeness_level not in (1, 2, 3) or defect_level not in (1, 2, 3):
            raise ValueError(f"Invalid rubric level for tomato {tomato_id}.")
        expected_grade = grade_from_levels(ripeness_level, defect_level)
        if str(row["grade"]).upper() != expected_grade:
            raise ValueError(
                f"Grade for tomato {tomato_id} conflicts with the rubric: "
                f"expected {expected_grade}, found {row['grade']!r}."
            )
        samples.append((path, row))
    if not samples:
        raise ValueError(f"No labelled primary photos found in {labels_path}.")
    return samples


def _evaluate_primary(
    data_root: Path,
    ripeness_model: Any,
    defect_model: Any,
    models_dir: Path,
) -> dict[str, Any]:
    from PIL import Image

    ripeness_metadata = json.loads(
        (models_dir / "ripeness_resnext101.json").read_text(encoding="utf-8")
    )
    defect_metadata = json.loads(
        (models_dir / "defect_resnext101.json").read_text(encoding="utf-8")
    )
    preprocess_input = get_preprocess_input()

    expected_grades: list[str] = []
    predicted_grades: list[str] = []
    ripeness_true: list[int] = []
    ripeness_predicted: list[int] = []
    defect_true: list[int] = []
    defect_predicted: list[int] = []
    primary_samples = _primary_test_samples(data_root)

    for path, labels in primary_samples:
        with Image.open(path) as source:
            image = source.convert("RGB").resize((INPUT_SIZE[1], INPUT_SIZE[0]))
            pixels = np.asarray(image, dtype=np.float32)
        processed = preprocess_input(np.expand_dims(pixels, axis=0))
        ripeness_probabilities = ripeness_model.predict(processed, verbose=0)[0]
        defect_probabilities = defect_model.predict(processed, verbose=0)[0]
        ripeness_classes = ripeness_metadata["class_names"]
        defect_classes = defect_metadata["class_names"]
        ripeness_class = ripeness_classes[int(np.argmax(ripeness_probabilities))]
        defect_class = defect_classes[int(np.argmax(defect_probabilities))]
        ripeness_level = RIPENESS_LEVELS[ripeness_class]
        defect_level = DEFECT_LEVELS[defect_class]
        expected_grades.append(str(labels["grade"]).upper())
        predicted_grades.append(grade_from_levels(ripeness_level, defect_level))
        ripeness_true.append(int(labels["tingkat_kematangan"]) - 1)
        ripeness_predicted.append(ripeness_classes.index(ripeness_class))
        if int(labels["tingkat_cacat"]) in (1, 3):
            defect_true.append(0 if int(labels["tingkat_cacat"]) == 1 else 1)
            defect_predicted.append(defect_classes.index(defect_class))

    grade_two_count = sum(
        1 for _, labels in primary_samples if int(labels["tingkat_cacat"]) == 2
    )
    return {
        "independent_primary_test_images": len(expected_grades),
        "grade": _grade_metrics(expected_grades, predicted_grades),
        "heads": {
            "ripeness": _classification_metrics(
                ripeness_true, ripeness_predicted, RIPENESS_CLASSES
            ),
            "defect": {
                **_classification_metrics(
                    defect_true, defect_predicted, DEFECT_CLASSES
                ),
                "evaluated_images_with_binary_ground_truth": len(defect_true),
                "excluded_level_2_ground_truth_images": grade_two_count,
            },
        },
    }


def _grade_metrics(expected: Sequence[str], predicted: Sequence[str]) -> dict[str, Any]:
    class_names = ("A", "B", "C")
    true_indices = [class_names.index(value) for value in expected]
    predicted_indices = [class_names.index(value) for value in predicted]
    return _classification_metrics(true_indices, predicted_indices, class_names)


def train(
    data_root: str | Path,
    models_dir: str | Path = "models",
    seed: int = 42,
    batch_size: int = 16,
    head_epochs: int = 8,
    fine_tune_epochs: int = 5,
    reuse_existing: bool = False,
) -> dict[str, Any]:
    """Train both heads and evaluate on secondary holdouts and primary photos."""
    if batch_size < 1:
        raise ValueError("batch_size must be greater than zero.")
    if head_epochs < 1:
        raise ValueError("head_epochs must be greater than zero.")
    if fine_tune_epochs < 0:
        raise ValueError("fine_tune_epochs cannot be negative.")
    root = Path(data_root)
    output_directory = Path(models_dir)
    secondary = root / "Sekunder"
    ripeness_root = secondary / "kematangan" / "Klasifikasi Kematangan Tomat"
    ripeness_directories = {
        "matang": ripeness_root / "Tomat Matang",
        "setengah_matang": ripeness_root / "Tomat Setengah Matang",
        "mentah": ripeness_root / "Tomat Mentah",
    }
    defect_root = secondary / "cacat" / "Two Classes"
    defect_directories = {
        "healthy": defect_root / "Healthy",
        "reject": defect_root / "Reject",
    }

    ripeness_samples = _collect_samples(ripeness_directories)
    defect_samples = _collect_samples(defect_directories)
    ripeness_model, ripeness_metadata = _train_task(
        "ripeness",
        ripeness_samples,
        RIPENESS_CLASSES,
        RIPENESS_LEVELS,
        output_directory / "ripeness_resnext101.keras",
        seed,
        batch_size,
        head_epochs,
        fine_tune_epochs,
        reuse_existing,
    )
    defect_model, defect_metadata = _train_task(
        "defect",
        defect_samples,
        DEFECT_CLASSES,
        DEFECT_LEVELS,
        output_directory / "defect_resnext101.keras",
        seed,
        batch_size,
        head_epochs,
        fine_tune_epochs,
        reuse_existing,
    )

    for filename, metadata in (
        ("ripeness_resnext101.json", ripeness_metadata),
        ("defect_resnext101.json", defect_metadata),
    ):
        (output_directory / filename).write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
    report = {
        "model": MODEL_NAME,
        "input_size": list(INPUT_SIZE),
        "seed": seed,
        "secondary": {
            "ripeness": ripeness_metadata["test_metrics"],
            "defect": defect_metadata["test_metrics"],
        },
        "primary_independent_test": _evaluate_primary(
            root, ripeness_model, defect_model, output_directory
        ),
    }
    (output_directory / "evaluation_resnext101.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-root", required=True, type=Path)
    parser.add_argument("--models-dir", type=Path, default=Path("models"))
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--batch-size", type=int, default=16)
    parser.add_argument("--head-epochs", type=int, default=8)
    parser.add_argument("--fine-tune-epochs", type=int, default=5)
    parser.add_argument(
        "--reuse-existing-models",
        action="store_true",
        help="Reuse existing .keras checkpoints instead of training those tasks again.",
    )
    args = parser.parse_args()
    results = train(
        args.data_root,
        args.models_dir,
        args.seed,
        args.batch_size,
        args.head_epochs,
        args.fine_tune_epochs,
        args.reuse_existing_models,
    )
    print(json.dumps(results, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
